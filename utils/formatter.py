from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF
import os


# =========================
# LOGO PATH
# =========================

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

LOGO_PATH = os.path.join(
    BASE_DIR,
    "assets",
    "LegalEase_logo.png"
)


# =========================
# TXT
# =========================

def save_as_txt(content, filename):

    with open(filename, "w", encoding="utf-8") as file:
        file.write(content)


# =========================
# DOCX
# =========================

def save_as_docx(content, filename):

    document = Document()

    # Add LegalEase logo
    if os.path.exists(LOGO_PATH):

        logo_paragraph = document.add_paragraph()

        logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = logo_paragraph.add_run()

        run.add_picture(
            LOGO_PATH,
            width=Inches(2.2)
        )

    # Add document content
    for line in content.split("\n"):

        if line.strip() == "":

            document.add_paragraph()

        elif line.isupper():

            paragraph = document.add_paragraph()

            run = paragraph.add_run(line)

            run.bold = True
            run.font.size = Pt(14)

        else:

            paragraph = document.add_paragraph(line)

            paragraph.paragraph_format.space_after = Pt(6)

    document.save(filename)


# =========================
# PDF
# =========================

class LegalEasePDF(FPDF):

    def header(self):

        # Add LegalEase logo
        if os.path.exists(LOGO_PATH):

            self.image(
                LOGO_PATH,
                x=80,
                y=10,
                w=50
            )

        self.ln(35)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Arial",
            size=8
        )

        self.cell(
            0,
            10,
            "LegalEase - Law Made Simple",
            align="C"
        )


def save_as_pdf(content, filename):

    # Handle unsupported characters
    content = content.encode(
        "latin-1",
        "replace"
    ).decode("latin-1")

    pdf = LegalEasePDF()

    pdf.set_auto_page_break(
        auto=True,
        margin=20
    )

    pdf.add_page()

    pdf.set_font(
        "Arial",
        size=12
    )

    for line in content.split("\n"):

        if line.strip() == "":

            pdf.ln(5)

        elif line.isupper():

            pdf.set_font(
                "Arial",
                style="B",
                size=14
            )

            pdf.multi_cell(
                0,
                8,
                line
            )

            pdf.set_font(
                "Arial",
                size=12
            )

        else:

            pdf.multi_cell(
                0,
                8,
                line
            )

    pdf.output(filename)


# =========================
# HTML PREVIEW
# =========================

def format_html_preview(content):

    html = ""

    for line in content.split("\n"):

        if line.strip() == "":

            html += "<br>"

        elif line.isupper() and len(line) < 80:

            html += f"<h3>{line}</h3>"

        else:

            html += f"<p>{line}</p>"

    return f"""
    <div style="
        background-color: #1e1e1e;
        color: white;
        padding: 25px;
        border-radius: 12px;
        max-height: 600px;
        overflow-y: auto;
        font-family: Arial, sans-serif;
        line-height: 1.6;
    ">
        {html}
    </div>
    """